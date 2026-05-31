#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

blank = prs.slide_layouts[6]  # blank
slide = prs.slides.add_slide(blank)

# ── Title ─────────────────────────────────────────────────────────────────────
tx = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10), Inches(1.2))
tf = tx.text_frame
tf.word_wrap = False
p  = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Experiments"
run.font.size = Pt(54)
run.font.bold = False
run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

# ── Body: Tasks label ─────────────────────────────────────────────────────────
# Each bullet row: left part (black) + right part "(fine tuning)" in light gray, right-aligned
# We use a table-like layout: two text boxes side by side per row

LIGHT_GRAY = RGBColor(0xAA, 0xAA, 0xAA)
BLACK      = RGBColor(0x11, 0x11, 0x11)

TOP   = Inches(3.0)
LEFT  = Inches(1.2)
W_L   = Inches(7.0)   # left column width
W_R   = Inches(3.5)   # right column (fine tuning) width
ROW_H = Inches(0.55)

# "Tasks:" label
tx = slide.shapes.add_textbox(LEFT, TOP, W_L, ROW_H)
tf = tx.text_frame
p  = tf.paragraphs[0]
run = p.add_run(); run.text = "Tasks:"; run.font.size = Pt(24); run.font.bold = True
run.font.color.rgb = BLACK

rows = [
    ("–   Next token-prediction",       ""),
    ("–   Factual recall",               "(fine tuning)"),
    ("–   Machine translation",          "(fine tuning)"),
]

for i, (label, annotation) in enumerate(rows):
    y = TOP + ROW_H + i * ROW_H

    # Left: bullet text
    tx_l = slide.shapes.add_textbox(LEFT, y, W_L, ROW_H)
    tf_l = tx_l.text_frame
    tf_l.word_wrap = False
    p = tf_l.paragraphs[0]
    run = p.add_run(); run.text = label; run.font.size = Pt(24)
    run.font.color.rgb = BLACK

    # Right: annotation in light gray, right-aligned
    if annotation:
        tx_r = slide.shapes.add_textbox(LEFT + W_L, y, W_R, ROW_H)
        tf_r = tx_r.text_frame
        tf_r.word_wrap = False
        p = tf_r.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run(); run.text = annotation; run.font.size = Pt(22)
        run.font.color.rgb = LIGHT_GRAY
        run.font.italic = True

out = "/Users/sttawm/Documents/NLP Presentation (1).pptx"
prs.save(out)
print(f"Saved → {out}")
