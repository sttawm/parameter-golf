#!/usr/bin/env python3
"""Rebuild slide 9 (Factual Recall Results) — two-column layout.

Left  (~3.5"): Test accuracy chart (square) + caption
Right (~5.9"): 24-run validation curves (hero) + caption with variance stats
Variance chart moved to appendix; its key numbers folded into right caption.
"""

import tempfile, os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

FONT       = "Helvetica Neue"
NEAR_BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BODY_GRAY  = RGBColor(0x44, 0x44, 0x44)
PASTEL_SEP = RGBColor(0xB8, 0xD4, 0xBE)

SRC_PPTX = "/Users/sttawm/Downloads/NLP Presentation (2).pptx"
OUT_PPTX = "/Users/sttawm/Downloads/NLP Presentation (2) - updated.pptx"

ACC_IMG    = "/tmp/s9_Google_Shape_250_p45.png"   # square accuracy chart
CURVES_IMG = "/tmp/s9_Google_Shape_251_p45.png"   # wide 24-run curves


def flatten(path):
    img = Image.open(path).convert("RGBA")
    bg  = Image.new("RGB", img.size, (255, 255, 255))
    bg.paste(img, mask=img.split()[3])
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    bg.save(tmp.name, "PNG")
    return tmp.name, img.size


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spPr = shape._element.find(qn("p:spPr"))
    if spPr is not None:
        el = spPr.find(qn("a:effectLst"))
        if el is None:
            el = etree.SubElement(spPr, qn("a:effectLst"))
        for child in list(el):
            el.remove(child)
    return shape


def add_caption(slide, text, left, top, width, height, size=12):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf  = box.text_frame; tf.word_wrap = True
    p   = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r   = p.add_run()
    r.text = text; r.font.size = Pt(size)
    r.font.name = FONT; r.font.color.rgb = BODY_GRAY


prs   = Presentation(SRC_PPTX)
slide = prs.slides[8]

# Clear existing shapes
sp_tree = slide.shapes._spTree
for sp in list(sp_tree)[2:]:
    sp_tree.remove(sp)

slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ── Title ─────────────────────────────────────────────────────────────────────
tb = slide.shapes.add_textbox(Inches(0.38), Inches(0.14), Inches(9.24), Inches(0.52))
tf = tb.text_frame; tf.word_wrap = False
r  = tf.paragraphs[0].add_run()
r.text = "Experiment #2: Factual Recall — Results"
r.font.size = Pt(24); r.font.bold = False
r.font.color.rgb = NEAR_BLACK; r.font.name = FONT

# ── Separator ─────────────────────────────────────────────────────────────────
add_rect(slide, 0.38, 0.68, 7.0, 0.022, PASTEL_SEP)

# ── Layout constants ──────────────────────────────────────────────────────────
MARGIN   = 0.28
GAP      = 0.35
IMG_TOP  = 0.88
CAP_H    = 0.60
USABLE_H = 5.62 - IMG_TOP - CAP_H - 0.25   # ~4.0" for images

# Left column: square accuracy chart — fix width, derive height
L_W = 3.3
acc_tmp, (aw, ah) = flatten(ACC_IMG)
L_H = L_W * ah / aw   # ~3.3" (nearly square)

# Right column: curves — takes remaining width
R_X = MARGIN + L_W + GAP
R_W = 10.0 - MARGIN - R_X
cur_tmp, (cw, ch) = flatten(CURVES_IMG)
R_H = R_W * ch / cw

# If curves too tall, scale both down proportionally
max_h = USABLE_H
if max(L_H, R_H) > max_h:
    scale = max_h / max(L_H, R_H)
    L_W *= scale; L_H *= scale
    R_W *= scale; R_H *= scale
    R_X  = MARGIN + L_W + GAP

# Vertically centre each image in the content area
def img_y(h):
    return IMG_TOP + (USABLE_H - h) / 2

# ── Left: accuracy chart ──────────────────────────────────────────────────────
slide.shapes.add_picture(acc_tmp, Inches(MARGIN), Inches(img_y(L_H)),
                         Inches(L_W), Inches(L_H))
os.unlink(acc_tmp)
add_caption(slide, "Final test accuracy (not significant)",
            MARGIN, img_y(L_H) + L_H + 0.10, L_W, CAP_H)

# ── Right: 24-run curves ──────────────────────────────────────────────────────
slide.shapes.add_picture(cur_tmp, Inches(R_X), Inches(img_y(R_H)),
                         Inches(R_W), Inches(R_H))
os.unlink(cur_tmp)
add_caption(slide,
            "Validation accuracy (λ=0 vs λ=1.0), averaged over 24 runs  ·  "
            "Lower variance (p=0.047); faster convergence (p=0.19, n.s.)",
            R_X, img_y(R_H) + R_H + 0.10, R_W, CAP_H)

prs.save(OUT_PPTX)
print(f"Saved → {OUT_PPTX}")
