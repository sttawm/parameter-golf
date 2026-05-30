#!/usr/bin/env python3
"""
1. Insert a new slide after slide 9 with the variance/convergence chart.
2. Add pastel separator to all [Name] placeholder slides.
"""

import copy, tempfile, os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn, nsmap
from lxml import etree

FONT       = "Helvetica Neue"
NEAR_BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BODY_GRAY  = RGBColor(0x44, 0x44, 0x44)
PASTEL_SEP = RGBColor(0xB8, 0xD4, 0xBE)

SRC  = "/Users/sttawm/Downloads/NLP Presentation (2) - updated.pptx"
OUT  = "/Users/sttawm/Downloads/NLP Presentation (2) - updated.pptx"
VAR_IMG = "/tmp/s9_Google_Shape_249_p45.png"


def flatten(path):
    img = Image.open(path).convert("RGBA")
    bg  = Image.new("RGB", img.size, (255, 255, 255))
    bg.paste(img, mask=img.split()[3])
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    bg.save(tmp.name, "PNG")
    return tmp.name, img.size


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spPr = shape._element.find(qn("p:spPr"))
    if spPr is not None:
        el = spPr.find(qn("a:effectLst"))
        if el is None:
            el = etree.SubElement(spPr, qn("a:effectLst"))
        for child in list(el):
            el.remove(child)
    return shape


prs = Presentation(SRC)

# ── 1. Add pastel separator to placeholder slides ─────────────────────────────
# Slides that have a PLACEHOLDER title (not our custom textboxes)
PLACEHOLDER_SLIDE_INDICES = [4, 6, 7, 9, 11, 12, 13]  # 0-indexed

for idx in PLACEHOLDER_SLIDE_INDICES:
    if idx >= len(prs.slides):
        continue
    slide = prs.slides[idx]
    # Skip if separator already present (check for our exact color)
    already = False
    for shape in slide.shapes:
        if shape.shape_type == 1:
            try:
                if shape.fill.fore_color.rgb == PASTEL_SEP:
                    already = True; break
            except: pass
    if already:
        continue
    add_rect(slide, 0.38, 1.14, 7.0, 0.022, PASTEL_SEP)
    print(f"  Added separator to slide {idx+1}")

# ── 2. Insert new slide after slide 9 (index 8) ──────────────────────────────
# Clone a blank layout from slide 9's layout
blank_layout = prs.slides[8].slide_layout

# Add slide at the end, then move it to position 9 (after current slide 9)
new_slide = prs.slides.add_slide(blank_layout)

# Move: remove from end, insert after index 8
xml_slides = prs.slides._sldIdLst
last = xml_slides[-1]
xml_slides.remove(last)
xml_slides.insert(9, last)

# Clear any inherited placeholder shapes
sp_tree = new_slide.shapes._spTree
for sp in list(sp_tree)[2:]:
    sp_tree.remove(sp)

new_slide.background.fill.solid()
new_slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Title
tb = new_slide.shapes.add_textbox(Inches(0.38), Inches(0.14), Inches(9.24), Inches(0.52))
tf = tb.text_frame; tf.word_wrap = False
r  = tf.paragraphs[0].add_run()
r.text = "Experiment #2: Factual Recall — Convergence & Variance"
r.font.size = Pt(24); r.font.bold = False
r.font.color.rgb = NEAR_BLACK; r.font.name = FONT

# Separator
add_rect(new_slide, 0.38, 0.68, 7.0, 0.022, PASTEL_SEP)

# Variance chart — centered, large
var_tmp, (vw, vh) = flatten(VAR_IMG)
IMG_W = 8.5
IMG_H = IMG_W * vh / vw   # ~5.26" → scale down
MAX_H = 3.9
if IMG_H > MAX_H:
    IMG_W = IMG_W * MAX_H / IMG_H
    IMG_H = MAX_H
img_x = (10.0 - IMG_W) / 2
img_y = 0.88

new_slide.shapes.add_picture(var_tmp, Inches(img_x), Inches(img_y),
                              Inches(IMG_W), Inches(IMG_H))
os.unlink(var_tmp)

# Caption
cap_box = new_slide.shapes.add_textbox(
    Inches(img_x), Inches(img_y + IMG_H + 0.12), Inches(IMG_W), Inches(0.55)
)
tf = cap_box.text_frame; tf.word_wrap = True
p  = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r  = p.add_run()
r.text = ("Lower variance (p=0.047)  ·  Faster convergence (p=0.19, n.s.)  ·  "
          "Each experiment run 24×, results averaged")
r.font.size = Pt(13); r.font.name = FONT; r.font.color.rgb = BODY_GRAY

prs.save(OUT)
print(f"Saved → {OUT}")
